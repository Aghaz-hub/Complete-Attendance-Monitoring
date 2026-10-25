"""
Attendance Monitoring Report Generator (v2)
- Multi-sheet workbook (First / Second / Third / Low / Final / Count)
- Fills only the selected monitoring sheet; other sheets left for manual entry
- Option: Separate Theory & Lab %  OR  Total subject %
- Supports .xlsx and .xls (legacy ERP exports)
"""

import streamlit as st
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter
from io import BytesIO
import re

from attendance_core import parse_erp_attendance

st.set_page_config(
    page_title="Attendance Monitoring Report Generator",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

THIN = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
HEADER_FONT = Font(bold=True, size=10)
TITLE_FONT = Font(bold=True, size=13)
HEADER_FILL = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
CRITICAL_FILL = PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")
LIGHT_RED = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

MONITOR_SHEETS = {
    "FIRST": "First Att Monitoring",
    "SECOND": "Second Att Monitoring",
    "THIRD": "Third Att Monitoring",
    "LOW": "Low Attendance Review",
    "FINAL": "Final Att Monitoring",
}


def load_excel_any(uploaded_file):
    """Load .xlsx via openpyxl; convert .xls via pandas/xlrd."""
    name = (uploaded_file.name or "").lower()
    data = uploaded_file.read()
    bio = BytesIO(data)

    if name.endswith(".xlsx"):
        return openpyxl.load_workbook(bio, data_only=True)

    try:
        return openpyxl.load_workbook(bio, data_only=True)
    except Exception:
        pass

    try:
        import pandas as pd
        bio.seek(0)
        xls = pd.ExcelFile(bio, engine="xlrd")
        wb = Workbook()
        wb.remove(wb.active)
        for sheet_name in xls.sheet_names:
            df = pd.read_excel(xls, sheet_name=sheet_name, header=None)
            ws = wb.create_sheet(title=str(sheet_name)[:31])
            for r_idx, row in df.iterrows():
                for c_idx, val in enumerate(row):
                    if pd.isna(val):
                        continue
                    ws.cell(int(r_idx) + 1, int(c_idx) + 1, val)
        return wb
    except Exception as e:
        raise ValueError(
            f"Could not read the Excel file ({e}). "
            "Please open it in Excel and Save As → Excel Workbook (*.xlsx), then upload again."
        )


def prepare_subjects(parsed, mode: str):
    """separate = Theory/Lab columns; total = merge into one % per subject."""
    subjects = parsed["subjects"]
    if mode == "separate":
        out = []
        for s in subjects:
            d = dict(s)
            d["merge_keys"] = [s["code"] or s["name"]]
            out.append(d)
        return out

    groups = {}
    order = []
    for s in subjects:
        base = s["name"]
        base = re.sub(r"\s*\(Theory\)\s*$", "", base, flags=re.I)
        base = re.sub(r"\s*\(Practical\)\s*$", "", base, flags=re.I)
        base = re.sub(r"\s+LAB\s*$", "", base, flags=re.I)
        base = base.strip()
        if base not in groups:
            groups[base] = []
            order.append(base)
        groups[base].append(s)

    merged = []
    for base in order:
        items = groups[base]
        codes = [it["code"] or it["name"] for it in items]
        merged.append({
            "name": base,
            "code": items[0].get("code") or base,
            "type": items[0].get("type") or "CORE",
            "faculty": items[0].get("faculty") or "",
            "component": "Total" if len(items) > 1 else items[0].get("component"),
            "merge_keys": codes,
            "pct_col": items[0].get("pct_col"),
            "d_col": items[0].get("d_col"),
        })
    return merged


def student_subject_pct(stu, subj, mode: str):
    keys = subj.get("merge_keys") or [subj["code"] or subj["name"]]
    if len(keys) == 1:
        info = stu["pcts"].get(keys[0], {})
        return info.get("pct"), info.get("d") or 0

    total_d = 0.0
    total_a = 0.0
    any_pct = None
    for key in keys:
        info = stu["pcts"].get(key, {})
        d = info.get("d") or 0
        p = info.get("pct")
        if d > 0 and p is not None:
            total_d += d
            total_a += d * (p / 100.0)
            any_pct = p
        elif p is not None and any_pct is None:
            any_pct = p
    if total_d > 0:
        return round(100.0 * total_a / total_d, 2), total_d
    return any_pct, 0


def build_workbook(parsed, monitoring, min_criteria, coordinator, mode):
    subjects = prepare_subjects(parsed, mode)
    students = parsed["students"]

    counts = []
    for stu in students:
        below = 0
        for s in subjects:
            pct, d = student_subject_pct(stu, s, mode)
            if d > 0 and pct is not None and pct < min_criteria:
                below += 1
        counts.append(below)

    wb = Workbook()
    wb.remove(wb.active)

    for period in ["FIRST", "SECOND", "THIRD", "LOW", "FINAL"]:
        ws = wb.create_sheet(MONITOR_SHEETS[period])
        fill_data = period == monitoring
        _write_monitoring_sheet(
            ws, parsed, subjects, students,
            counts if fill_data else None,
            period, min_criteria, coordinator, mode, fill_data,
        )

    ws_c = wb.create_sheet("Monitoring-wise Count Sheet")
    _write_count_sheet(ws_c, students, counts, monitoring, min_criteria)
    return wb


def _write_monitoring_sheet(
    ws, parsed, subjects, students, counts, period, min_criteria,
    coordinator, mode, fill_data,
):
    meta = parsed["meta"]
    label = {
        "FIRST": "FIRST", "SECOND": "SECOND", "THIRD": "THIRD",
        "LOW": "LOW ATTENDANCE REVIEW", "FINAL": "FINAL",
    }.get(period, period)

    ws.merge_cells("A2:H2")
    ws["A2"] = f"DEPARTMENT OF CST\nJAN-MAY 2026\n{label} ATTENDANCE MONITORING REPORT"
    ws["A2"].font = TITLE_FONT
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 48

    ws["A5"] = f"Branch: {meta.get('branch') or 'MRU-School of Engineering'}"
    ws["A6"] = f"Department: {meta.get('department') or 'Bachelor of Technology in Computer Science and Engineering'}"
    ws["A7"] = f"Class Name: {meta.get('class_name') or ''}"
    ws["A8"] = f"Division: {meta.get('division') or 'All'}"
    ws["A9"] = f"Date: {meta.get('date_range') or ''}"
    ws["A10"] = f"Program Coordinator: {coordinator or ''}"
    for r in range(5, 11):
        ws.cell(r, 1).font = Font(size=10)

    n_subj = len(subjects)
    header_row, code_row, type_row, faculty_row, data_start = 12, 13, 14, 15, 16
    overall_col = 4 + n_subj
    count_col = overall_col + 2
    crit_col = overall_col + 3

    for i, h in enumerate(["Sr No.", "Roll No", "Student Name"], 1):
        cell = ws.cell(header_row, i, h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN
        cell.alignment = CENTER

    for i, s in enumerate(subjects):
        col = 4 + i
        cell = ws.cell(header_row, col, s["name"])
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN
        cell.alignment = CENTER
        ws.cell(code_row, col, s.get("code") or "").border = THIN
        ws.cell(code_row, col).font = Font(size=8)
        ws.cell(code_row, col).alignment = CENTER
        typ = s.get("type") or "CORE"
        if s.get("component") == "Theory":
            typ = f"{typ} (Theory)"
        elif s.get("component") == "Practical":
            typ = f"{typ} (Lab)"
        elif s.get("component") == "Total":
            typ = f"{typ} (Total)"
        ws.cell(type_row, col, typ).border = THIN
        ws.cell(type_row, col).font = Font(size=8)
        ws.cell(faculty_row, col, s.get("faculty") or "").border = THIN
        ws.cell(faculty_row, col).font = Font(size=8)

    for col, text in [
        (overall_col, "Overall %age of all subjects from ERP report"),
        (overall_col + 1, "Roll No"),
        (count_col, "Count of Courses with attendance below minimum attendance criteria"),
        (crit_col, "Whether Critical"),
    ]:
        cell = ws.cell(type_row, col, text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN
        cell.alignment = CENTER

    ws.cell(header_row, crit_col + 1, f"Minimum Attendance Criteria: {min_criteria}%")
    ws.cell(header_row, crit_col + 1).font = Font(size=9, italic=True)
    mode_label = "Separate Theory & Lab %" if mode == "separate" else "Total subject %"
    ws.cell(11, 1, f"Attendance display mode: {mode_label}")
    ws.cell(11, 1).font = Font(size=9, italic=True, color="666666")

    for idx, stu in enumerate(students):
        row = data_start + idx
        ws.cell(row, 1, idx + 1).border = THIN
        ws.cell(row, 1).alignment = CENTER
        ws.cell(row, 2, stu["roll"]).border = THIN
        ws.cell(row, 2).alignment = CENTER
        ws.cell(row, 3, stu["name"]).border = THIN

        if fill_data:
            for i, s in enumerate(subjects):
                col = 4 + i
                pct, d = student_subject_pct(stu, s, mode)
                cell = ws.cell(row, col)
                cell.border = THIN
                cell.alignment = CENTER
                if pct is not None:
                    cell.value = round(pct, 2) if pct != int(pct) else int(pct)
                    if d > 0 and pct < min_criteria:
                        cell.fill = LIGHT_RED

            cell = ws.cell(row, overall_col, stu.get("total_pct"))
            cell.border = THIN
            cell.alignment = CENTER
            ws.cell(row, overall_col + 1, stu["roll"]).border = THIN

            below = counts[idx]
            cell = ws.cell(row, count_col, below)
            cell.border = THIN
            cell.alignment = CENTER
            if below > 0:
                cell.fill = LIGHT_RED

            is_crit = "CRITICAL" if below > 2 else ""
            cell = ws.cell(row, crit_col, is_crit)
            cell.border = THIN
            cell.alignment = CENTER
            if is_crit:
                cell.fill = CRITICAL_FILL
                cell.font = Font(bold=True, size=10, color="FFFFFF")
        else:
            for i in range(n_subj):
                ws.cell(row, 4 + i).border = THIN
            for col in (overall_col, overall_col + 1, count_col, crit_col):
                ws.cell(row, col).border = THIN

    last_data = data_start + len(students) - 1
    n = len(students)

    # ── Course-wise compilation rows (match departmental sample) ──
    # Placed after student block; formulas/values per subject column
    compile_start = last_data + 2
    labels = [
        ("Number of Students in course", None),
        (f"Number of students below {min_criteria}%", min_criteria),
        ("Number of students below 70%", 70),
        ("Number of students below 65%", 65),
        ("Number of students below 60%", 60),
        (f"%age of students below {min_criteria}%", "pct"),
    ]

    if fill_data and students and n_subj > 0:
        # Precompute per-subject stats
        # For each subject: enrolled (d>0), below thresholds
        subj_stats = []  # list of dicts
        for s in subjects:
            enrolled = 0
            b75 = b70 = b65 = b60 = 0
            for stu in students:
                pct, d = student_subject_pct(stu, s, mode)
                if d > 0 and pct is not None:
                    enrolled += 1
                    if pct < min_criteria:
                        b75 += 1
                    if pct < 70:
                        b70 += 1
                    if pct < 65:
                        b65 += 1
                    if pct < 60:
                        b60 += 1
            pct_below = round(100.0 * b75 / enrolled, 1) if enrolled else 0
            subj_stats.append({
                "enrolled": enrolled, "b75": b75, "b70": b70,
                "b65": b65, "b60": b60, "pct_below": pct_below,
            })

        row_enrolled = compile_start
        row_b75 = compile_start + 1
        row_b70 = compile_start + 2
        row_b65 = compile_start + 3
        row_b60 = compile_start + 4
        row_pct = compile_start + 5

        ws.cell(row_enrolled, 3, "Number of Students in course").font = Font(bold=True, size=9)
        ws.cell(row_b75, 3, f"Number of students below {min_criteria}%").font = Font(bold=True, size=9)
        ws.cell(row_b70, 3, "Number of students below 70%").font = Font(bold=True, size=9)
        ws.cell(row_b65, 3, "Number of students below 65%").font = Font(bold=True, size=9)
        ws.cell(row_b60, 3, "Number of students below 60%").font = Font(bold=True, size=9)
        ws.cell(row_pct, 3, f"%age of students below {min_criteria}%").font = Font(bold=True, size=9)

        for i, st in enumerate(subj_stats):
            col = 4 + i
            for r, key in [
                (row_enrolled, "enrolled"), (row_b75, "b75"), (row_b70, "b70"),
                (row_b65, "b65"), (row_b60, "b60"), (row_pct, "pct_below"),
            ]:
                cell = ws.cell(r, col, st[key])
                cell.border = THIN
                cell.alignment = CENTER
                cell.font = Font(size=9)
                if key == "pct_below" and st[key] >= 50:
                    cell.fill = LIGHT_RED

        # Overall / critical summary
        sum_row = row_pct + 2
        below_any = sum(1 for c in counts if c > 0)
        crit_n = sum(1 for c in counts if c > 2)
        ws.cell(sum_row, 1, f"Total students: {n}")
        ws.cell(sum_row + 1, 1, f"Students with at least 1 course below {min_criteria}%: {below_any}")
        ws.cell(sum_row + 2, 1, f"CRITICAL students (more than 2 courses below): {crit_n}")
        for r in range(sum_row, sum_row + 3):
            ws.cell(r, 1).font = Font(bold=True, size=10)

        # Chart data — same format as departmental sample:
        # Category row: course names (X-axis labels under bars)
        # Series rows: First / Second / Third Att Monitoring (% below criteria)
        # Only selected period is filled; other series = 0 (structure for comparison)
        cat_row = compile_start + 7
        series_first = cat_row + 1
        series_second = cat_row + 2
        series_third = cat_row + 3

        ws.cell(cat_row, 3, f"%age of students below {min_criteria}%").font = Font(bold=True, size=9)
        for i, s in enumerate(subjects):
            col = 4 + i
            ws.cell(cat_row, col, s["name"]).font = Font(size=7)
            ws.cell(cat_row, col).alignment = CENTER

        period_series = [
            ("FIRST", series_first, "First Att Monitoring"),
            ("SECOND", series_second, "Second Att Monitoring"),
            ("THIRD", series_third, "Third Att Monitoring"),
        ]
        for pkey, r, plabel in period_series:
            ws.cell(r, 3, plabel).font = Font(bold=True, size=9)
            for i, st in enumerate(subj_stats):
                col = 4 + i
                val = st["pct_below"] if pkey == period else 0
                cell = ws.cell(r, col, val)
                cell.border = THIN
                cell.alignment = CENTER
                cell.font = Font(size=9)

        # Grouped clustered bar chart (sample format)
        if n_subj > 0:
            chart = BarChart()
            chart.type = "col"
            chart.grouping = "clustered"
            chart.style = 10
            chart.title = f"{label.title()} Attendance Monitoring"
            chart.y_axis.title = f"%age of students below {min_criteria}%"
            chart.x_axis.title = "Course"

            data_ref = Reference(
                ws,
                min_col=3,
                max_col=3 + n_subj,
                min_row=series_first,
                max_row=series_third,
            )
            cats = Reference(
                ws,
                min_col=4,
                max_col=3 + n_subj,
                min_row=cat_row,
            )
            chart.add_data(data_ref, from_rows=True, titles_from_data=True)
            chart.set_categories(cats)
            chart.shape = 4
            chart.width = 22
            chart.height = 12
            if chart.legend:
                chart.legend.position = "b"
            ws.add_chart(chart, f"A{series_third + 3}")

        obs_row = series_third + 18
    else:
        # Empty structure: still reserve compile label rows
        for i, (lab, _) in enumerate(labels):
            ws.cell(compile_start + i, 3, lab).font = Font(bold=True, size=9)
            for j in range(n_subj):
                ws.cell(compile_start + i, 4 + j).border = THIN
        obs_row = compile_start + len(labels) + 3

    ws.cell(obs_row, 1, "Observations:")
    ws.cell(obs_row + 1, 1, "Action Taken:")
    ws.cell(obs_row + 2, 1, "Action Taken for Critical Cases:")
    for r in range(obs_row, obs_row + 3):
        ws.cell(r, 1).font = Font(bold=True, size=10)

    sig_row = obs_row + 6
    ws.cell(sig_row, 1, "Signature of Mentor")
    ws.cell(sig_row + 1, 1, "Mentor Name:")
    ws.cell(sig_row, 5, "HoD")
    ws.cell(sig_row + 3, 1, "Dean of School")
    ws.cell(sig_row + 3, 5, "Associate Dean (Academics)")
    ws.cell(sig_row + 5, 5, "Dean (Academics)")
    for r, c in [(sig_row, 1), (sig_row + 1, 1), (sig_row, 5),
                 (sig_row + 3, 1), (sig_row + 3, 5), (sig_row + 5, 5)]:
        ws.cell(r, c).font = Font(bold=True, size=10)

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 28
    for i in range(n_subj):
        ws.column_dimensions[get_column_letter(4 + i)].width = 12
    ws.column_dimensions[get_column_letter(overall_col)].width = 12
    ws.column_dimensions[get_column_letter(count_col)].width = 12
    ws.column_dimensions[get_column_letter(crit_col)].width = 12


def _write_count_sheet(ws, students, counts, monitoring, min_criteria):
    ws["A2"] = "DEPARTMENT OF CST — MONITORING-WISE REPORT OF STUDENTS WITH ATTENDANCE BELOW MINIMUM CRITERIA"
    ws["A2"].font = TITLE_FONT
    ws.merge_cells("A2:G2")
    ws["D3"] = f"Count of courses with attendance below {min_criteria}%"
    ws["D3"].font = Font(bold=True, size=10)

    headers = [
        "S. No.", "Roll No", "Student Name",
        "First monitoring", "Second Monitoring", "Third monitoring",
        "Low Attendance Review", "Final Monitoring",
    ]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(4, i, h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN
        cell.alignment = CENTER

    col_map = {"FIRST": 4, "SECOND": 5, "THIRD": 6, "LOW": 7, "FINAL": 8}
    fill_col = col_map.get(monitoring, 4)

    for idx, stu in enumerate(students):
        row = 5 + idx
        ws.cell(row, 1, idx + 1).border = THIN
        ws.cell(row, 2, stu["roll"]).border = THIN
        ws.cell(row, 3, stu["name"]).border = THIN
        for c in range(4, 9):
            cell = ws.cell(row, c)
            cell.border = THIN
            cell.alignment = CENTER
            if c == fill_col:
                cell.value = counts[idx]
                if counts[idx] > 2:
                    cell.fill = CRITICAL_FILL
                    cell.font = Font(bold=True, color="FFFFFF")
                elif counts[idx] > 0:
                    cell.fill = LIGHT_RED

    for letter, w in [("A", 10), ("B", 16), ("C", 24)]:
        ws.column_dimensions[letter].width = w
    for c in range(4, 9):
        ws.column_dimensions[get_column_letter(c)].width = 14

    note = 5 + len(students) + 2
    ws.cell(note, 1, f"Note: Only '{monitoring}' column is auto-filled from the uploaded ERP file.")
    ws.cell(note + 1, 1, "Other monitoring columns are left blank for manual entry (or generate those reports separately and copy values).")
    ws.cell(note, 1).font = Font(size=9, italic=True, color="666666")
    ws.cell(note + 1, 1).font = Font(size=9, italic=True, color="666666")


# ── UI ──
st.markdown("## 📊 Attendance Monitoring Report Generator")
st.caption("Upload ERP attendance → choose monitoring period & display mode → download multi-sheet departmental report")

with st.sidebar:
    st.header("⚙️ Settings")
    monitoring = st.radio(
        "Monitoring Period",
        options=["FIRST", "SECOND", "THIRD"],
        index=1,
        help="Only this sheet is filled from ERP. Other sheets stay blank for manual entry.",
    )
    mode = st.radio(
        "Attendance display mode",
        options=["separate", "total"],
        format_func=lambda x: (
            "Separate Theory % and Lab %"
            if x == "separate"
            else "Total % per subject (e.g. 7th sem style)"
        ),
        index=0,
        help="Separate: Theory and Lab as different columns. Total: one combined % per subject.",
    )
    min_criteria = st.number_input(
        "Minimum Attendance Criteria (%)",
        min_value=50, max_value=100, value=75, step=1,
    )
    coordinator = st.text_input("Program Coordinator", placeholder="e.g. Mr. / Ms. Name")
    st.divider()
    st.caption("CRITICAL when more than 2 courses are below criteria.")
    st.caption("Supports .xlsx and .xls (legacy ERP).")
    st.caption("Output sheets: First / Second / Third / Low / Final / Count.")

uploaded = st.file_uploader(
    "Upload Attendance Monitoring Excel file",
    type=["xlsx", "xls"],
    help="ERP export (.xlsx preferred; .xls also supported)",
)

if not uploaded:
    st.info(
        "👆 Upload an ERP attendance Excel. "
        "Choose **FIRST / SECOND / THIRD** — only that sheet is filled; "
        "other monitoring sheets remain empty for you to complete later."
    )
    with st.expander("What's new"):
        st.markdown(
            """
            - **Multi-sheet report** (First / Second / Third / Low / Final / Count)
            - **Only selected monitoring** auto-filled from ERP
            - Other periods left **blank for manual entry**
            - **Display mode**: Separate Theory & Lab **or** Total % per subject
            - **CRITICAL** if more than 2 courses below criteria
            - Supports **.xls** and **.xlsx**
            - Bar chart + Observations / Action Taken / signatures
            """
        )
    st.stop()

try:
    wb_in = load_excel_any(uploaded)
    parsed = parse_erp_attendance(wb_in)
except Exception as e:
    st.error(str(e))
    st.stop()

n_stu = len(parsed["students"])
subjects_preview = prepare_subjects(parsed, mode)
n_subj = len(subjects_preview)

if n_stu == 0:
    st.warning("No students found. Check that the file is a standard ERP Attendance Monitoring export.")
    st.stop()

st.success(
    f"Loaded **{n_stu}** students · **{n_subj}** subject columns "
    f"({'Theory+Lab separate' if mode == 'separate' else 'Total % mode'}) · "
    f"will fill **{monitoring}** sheet only"
)

crit_preview = []
for stu in parsed["students"]:
    below = 0
    for s in subjects_preview:
        pct, d = student_subject_pct(stu, s, mode)
        if d > 0 and pct is not None and pct < min_criteria:
            below += 1
    if below > 2:
        crit_preview.append({
            "Roll No": stu["roll"],
            "Name": stu["name"],
            "Courses below": below,
            "Overall %": stu.get("total_pct"),
        })

st.markdown(f"#### Critical students preview ({len(crit_preview)} / {n_stu})")
if crit_preview:
    st.dataframe(crit_preview, use_container_width=True, hide_index=True)
else:
    st.caption("No student has more than 2 courses below criteria.")

if st.button("Generate Formatted Report", type="primary", use_container_width=True):
    with st.spinner("Building multi-sheet report..."):
        wb_out = build_workbook(parsed, monitoring, min_criteria, coordinator, mode)
        buf = BytesIO()
        wb_out.save(buf)
        buf.seek(0)
        class_slug = re.sub(r"[^\w]+", "_", parsed["meta"].get("class_name") or "Class")
        fname = f"{class_slug}_{monitoring.title()}_Att_Monitoring.xlsx"
        st.download_button(
            "⬇️ Download multi-sheet report",
            data=buf.getvalue(),
            file_name=fname,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )
        st.success(
            f"✅ Report ready — **{monitoring}** sheet filled · "
            f"other monitoring sheets blank for manual entry · "
            f"{len(crit_preview)} CRITICAL"
        )
